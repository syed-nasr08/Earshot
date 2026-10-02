"""Writes questions and answers into a .docx file. Thread-safe."""
import os
import threading

from docx import Document

from config import DOC_PATH


class DocWriter:
    def __init__(self):
        self._lock = threading.Lock()
        self.path = DOC_PATH
        if not os.path.exists(self.path):
            self._new_doc().save(self.path)

    @staticmethod
    def _new_doc() -> Document:
        doc = Document()
        doc.add_heading("Interview Answers", level=0)
        return doc

    def append(self, question: str, answer: str):
        with self._lock:
            doc = Document(self.path)
            q = doc.add_paragraph()
            run = q.add_run(f"Q: {question}")
            run.bold = True
            doc.add_paragraph(f"A: {answer}")
            doc.add_paragraph("")  # spacer between entries
            doc.save(self.path)

    def clear(self):
        """Wipe the entire document and start fresh (bound to F9)."""
        with self._lock:
            self._new_doc().save(self.path)
        print("\n[F9] Document cleared — fresh session.\n")
